import csv
import hashlib
import os
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from dateutil.parser import parse as parse_date

class BankTransactionImporter:
    """Imports and normalizes bank transactions from local CSV/Excel exports."""

    DEFAULT_COLUMN_ALIASES = {
        "date": ["date", "datum", "transaction_date", "boekdatum", "rentedatum"],
        "amount": ["amount", "bedrag", "transactiebedrag", "af_bij", "af/bij"],
        "description": ["description", "omschrijving", "mededelingen", "details", "name_description"],
        "counterparty": ["counterparty", "tegenpartij", "naam_tegenpartij", "rekeninghouder", "name"],
        "currency": ["currency", "valuta", "muntsoort"],
        "reference": ["reference", "referentie", "kenmerk", "payment_reference"],
        "account": ["account", "rekening", "iban", "account_number"],
    }

    def __init__(self, config: Dict[str, Any]):
        self.config = config or {}
        self.bank_import_dir = Path(self.config.get("bank_import_dir", "data/bank_exports"))
        self.default_currency = self.config.get("default_currency", "EUR")
        self.column_aliases = dict(self.DEFAULT_COLUMN_ALIASES)
        self.column_aliases.update(self.config.get("bank_column_aliases", {}))

    def import_transactions(self) -> List[Dict[str, Any]]:
        self.bank_import_dir.mkdir(parents=True, exist_ok=True)
        transactions: List[Dict[str, Any]] = []
        for file_path in sorted(self.bank_import_dir.iterdir()):
            if not file_path.is_file():
                continue
            if file_path.suffix.lower() not in {".csv", ".xlsx"}:
                continue
            transactions.extend(self.import_file(file_path))
        return transactions

    def import_file(self, file_path: Any) -> List[Dict[str, Any]]:
        path = Path(file_path)
        rows = self._read_rows(path)
        normalized: List[Dict[str, Any]] = []
        for row_index, row in enumerate(rows):
            transaction = self._normalize_row(row, path, row_index)
            if transaction:
                normalized.append(transaction)
        return normalized

    def _read_rows(self, path: Path) -> List[Dict[str, Any]]:
        suffix = path.suffix.lower()
        if suffix == ".csv":
            return self._read_csv_rows(path)
        if suffix == ".xlsx":
            return self._read_xlsx(path)
        return []

    def _read_xlsx(self, path: Path) -> List[Dict[str, Any]]:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                return []
            normalized_headers = [self._normalize_column_name(value) for value in headers]
            result = []
            for values in rows:
                if not any(not self._is_empty(value) for value in values):
                    continue
                result.append({
                    header: value
                    for header, value in zip(normalized_headers, values)
                    if header
                })
            return result
        finally:
            workbook.close()

    def _read_csv_rows(self, path: Path) -> List[Dict[str, Any]]:
        for encoding in ["utf-8-sig", "latin-1"]:
            for separator in [",", ";", "\t"]:
                try:
                    with open(path, newline="", encoding=encoding) as handle:
                        reader = csv.DictReader(handle, delimiter=separator)
                        if not reader.fieldnames or len(reader.fieldnames) <= 1:
                            continue
                        return [
                            {
                                self._normalize_column_name(key): value
                                for key, value in row.items()
                                if key is not None
                            }
                            for row in reader
                            if any(not self._is_empty(value) for value in row.values())
                        ]
                except UnicodeDecodeError:
                    continue
                except OSError:
                    return []
        return []

    def _normalize_row(self, row: Dict[str, Any], source_path: Path, row_index: int) -> Optional[Dict[str, Any]]:
        date_value = self._first_present(row, "date")
        amount_value = self._first_present(row, "amount")
        description = self._first_present(row, "description") or ""
        counterparty = self._first_present(row, "counterparty") or ""
        currency = self._first_present(row, "currency") or self.default_currency
        reference = self._first_present(row, "reference") or ""
        account = self._first_present(row, "account") or ""

        parsed_date = self._parse_date(date_value)
        parsed_amount = self._parse_amount(amount_value)
        if not parsed_date or parsed_amount is None:
            return None

        raw_payload = {
            "source_file": str(source_path),
            "row_index": row_index,
            "row": row,
        }
        transaction_id = self._build_transaction_id(parsed_date, parsed_amount, description, counterparty, reference, account, row_index)
        return {
            "id": transaction_id,
            "source": "bank_export",
            "date": parsed_date,
            "transaction_date": parsed_date,
            "amount": parsed_amount,
            "currency": str(currency).strip() or self.default_currency,
            "description": str(description).strip(),
            "counterparty": str(counterparty).strip(),
            "reference": str(reference).strip(),
            "account": str(account).strip(),
            "metadata": raw_payload,
        }

    def _first_present(self, row: Dict[str, Any], canonical_name: str) -> Any:
        for alias in self.column_aliases.get(canonical_name, []):
            normalized_alias = self._normalize_column_name(alias)
            if normalized_alias in row and not self._is_empty(row[normalized_alias]):
                return row[normalized_alias]
        return None

    @staticmethod
    def _normalize_column_name(value: Any) -> str:
        return str(value).strip().lower().replace(" ", "_").replace("-", "_")

    @staticmethod
    def _is_empty(value: Any) -> bool:
        if value is None:
            return True
        if isinstance(value, float) and value != value:
            return True
        return str(value).strip() == ""

    @staticmethod
    def _parse_amount(value: Any) -> Optional[float]:
        if value is None:
            return None
        if isinstance(value, (int, float)):
            return float(value)
        text = str(value).strip().replace("€", "").replace(" ", "")
        if not text:
            return None
        if "," in text and "." in text:
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        else:
            text = text.replace(",", ".")
        try:
            return float(text)
        except ValueError:
            return None

    @staticmethod
    def _parse_date(value: Any) -> Optional[str]:
        if value is None:
            return None
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        if hasattr(value, "date"):
            return value.date().isoformat()
        text = str(value).strip()
        try:
            return datetime.fromisoformat(text.replace("Z", "+00:00")).date().isoformat()
        except ValueError:
            pass
        for fmt in ["%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d", "%d.%m.%Y", "%Y%m%d"]:
            try:
                return datetime.strptime(text, fmt).date().isoformat()
            except ValueError:
                continue
        try:
            return parse_date(text, dayfirst=True, fuzzy=False).date().isoformat()
        except (OverflowError, TypeError, ValueError):
            pass
        return None

    @staticmethod
    def _build_transaction_id(date: str, amount: float, description: str, counterparty: str, reference: str, account: str, row_index: int) -> str:
        source = "|".join([date, f"{amount:.2f}", description, counterparty, reference, account, str(row_index)])
        return "bank_" + hashlib.sha256(source.encode("utf-8")).hexdigest()
